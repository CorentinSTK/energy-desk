"""European Commission Weekly Oil Bulletin: consumer prices of petroleum products by member state.

`source_ref` is "<with_tax|wo_tax>|<country code>|<euro95|diesel>", e.g. "with_tax|DE|diesel";
"EU" is the EU average and "EUR" the euro-area average, both computed by the Commission.

Checked live on 2026-10-30 (file "Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx", 4.5 MB):
- One workbook holds the full history; the sheets used are "Prices with taxes" and "Prices wo
  taxes". Row 0 holds column names such as "DE_price_with_tax_euro95", row 2 the units ("1000 l"),
  then one row per bulletin date (a Monday, newest first) from 3 January 2005, then a "Notes:" row.
- Prices are in euros per 1,000 litres for every country, including those outside the euro area
  (converted by the Commission). Some weeks are skipped (holidays): they are gaps, not filled.
- Columns are found by name, never by position. A changed unit or a non-numeric price cell fails
  the run loudly rather than storing a misread value.
- Terms: the page's legal notice (commission.europa.eu/legal-notice) licenses EU-owned content under
  CC BY 4.0 (Commission Decision 2011/833/EU): reuse allowed with credit, changes indicated.
"""

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import httpx
from openpyxl import load_workbook

from worker.catalog import SeriesSpec
from worker.http_client import get
from worker.sources.base import Source
from worker.validate import Observation

URL = (
    "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)
SHEETS = {"with_tax": "Prices with taxes", "wo_tax": "Prices wo taxes"}
HEADER_ROW = 0
UNITS_ROW = 2
FIRST_DATA_ROW = 3

# column name → [(bulletin date, value or None)]
Columns = dict[str, list[tuple[date, Any]]]


class BulletinFormatError(Exception):
    pass


def column_name(kind: str, country: str, product: str) -> str:
    return f"{country}_price_{kind}_{product}"


def parse_sheet(rows: list[tuple[Any, ...]]) -> tuple[Columns, dict[str, Any]]:
    """Named columns of one sheet, with the unit stated under each column."""
    if len(rows) <= FIRST_DATA_ROW:
        raise BulletinFormatError(f"sheet has {len(rows)} rows, expected a header and data")
    header = rows[HEADER_ROW]
    units = rows[UNITS_ROW]
    dated = [r for r in rows[FIRST_DATA_ROW:] if isinstance(r[0], datetime)]
    columns: Columns = {}
    unit_of: dict[str, Any] = {}
    for i, name in enumerate(header):
        if not isinstance(name, str) or "_price_" not in name:
            continue
        unit_of[name] = units[i] if i < len(units) else None
        columns[name] = [(r[0].date(), r[i] if i < len(r) else None) for r in dated]
    return columns, unit_of


def observations_for(
    columns: Columns, unit_of: dict[str, Any], spec: SeriesSpec, since: date | None
) -> list[Observation]:
    kind, country, product = spec.source_ref.split("|")
    name = column_name(kind, country, product)
    if name not in columns:
        raise BulletinFormatError(f"{spec.code}: column {name!r} not found")
    if spec.source_unit is not None and unit_of[name] != spec.source_unit:
        raise BulletinFormatError(
            f"{spec.code}: expected unit {spec.source_unit!r}, got {unit_of[name]!r}"
        )
    observations = []
    for day, value in columns[name]:
        if value is None or (since is not None and day < since):
            continue
        if isinstance(value, bool) or not isinstance(value, int | float):
            raise BulletinFormatError(f"{spec.code}: non-numeric price {value!r} on {day}")
        ts = datetime(day.year, day.month, day.day, tzinfo=UTC)
        observations.append(Observation(spec.code, ts, float(value)))
    observations.sort(key=lambda o: o.ts)
    return observations


def read_workbook(content: bytes) -> dict[str, list[tuple[Any, ...]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        return {
            kind: list(workbook[sheet].iter_rows(values_only=True))
            for kind, sheet in SHEETS.items()
        }
    finally:
        workbook.close()


class EcWobSource(Source):
    name = "ec_wob"
    request_interval = 0.0  # one download per run serves every series: no pause between series

    def __init__(self) -> None:
        self._sheets: dict[str, tuple[Columns, dict[str, Any]]] | None = None

    def _load(self, client: httpx.Client) -> dict[str, tuple[Columns, dict[str, Any]]]:
        """Download and parse the workbook once per run (every series reads the same file)."""
        if self._sheets is None:
            raw = read_workbook(get(client, URL).content)
            self._sheets = {kind: parse_sheet(rows) for kind, rows in raw.items()}
        return self._sheets

    def fetch(
        self, client: httpx.Client, spec: SeriesSpec, since: date | None
    ) -> list[Observation]:
        kind = spec.source_ref.split("|")[0]
        columns, unit_of = self._load(client)[kind]
        return observations_for(columns, unit_of, spec, since)
